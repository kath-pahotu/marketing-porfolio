"""Generate matching, two-page ATS-readable CVs from the verified content."""
from pathlib import Path
import shutil
from html import escape
import argparse
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from revamp_content import PROFILE as P, TRACKS, EXPERIENCE, CASES, SKILLS, CV_PROJECTS, EDUCATION

FONT=Path('C:/Windows/Fonts')
pdfmetrics.registerFont(TTFont('CVBody',str(FONT/'arial.ttf')))
pdfmetrics.registerFont(TTFont('CVBold',str(FONT/'arialbd.ttf')))
pdfmetrics.registerFontFamily('CVBody',normal='CVBody',bold='CVBold',italic='CVBody',boldItalic='CVBold')
STYLES={}
for name,size,leading,opts in [
 ('name',25,29,dict(fontName='CVBold',spaceAfter=4)),
 ('title',12,17,dict(fontName='CVBold',spaceAfter=7)),
 ('contact',9,13,dict(spaceAfter=2)),
 ('heading',12,16,dict(fontName='CVBold',spaceBefore=13,spaceAfter=7,keepWithNext=True)),
 ('role',10.8,14,dict(fontName='CVBold',spaceBefore=8,spaceAfter=3,keepWithNext=True)),
 ('meta',9.2,12.6,dict(textColor=colors.HexColor('#415366'),spaceAfter=5,keepWithNext=True)),
 ('body',10.5,14.6,dict(spaceAfter=6)),
 ('bullet',10.5,14.6,dict(spaceAfter=5,leftIndent=10,firstLineIndent=-10))]:
    STYLES[name]=ParagraphStyle(name,fontName=opts.pop('fontName','CVBody'),fontSize=size,leading=leading,textColor=opts.pop('textColor',colors.black),**opts)

def blocks(track):
    t=TRACKS[track]; out=[]
    def add(kind,text):out.append((kind,text))
    add('name',P['name']);add('title',t['role'])
    add('contact',P['location']+' | '+P['phone']+' | '+P['email'])
    add('links',[(t['site'].replace('https://',''),t['site']),('LinkedIn',P['linkedin']),('GitHub',P['github'])])
    add('heading','Professional summary');add('body',t['summary'])
    add('heading','Core capabilities')
    for label,text in SKILLS[track]:add('skill',(label,text))
    add('heading','Professional experience')
    jobs=EXPERIENCE if track=='marketing' else [EXPERIENCE[1],EXPERIENCE[0],EXPERIENCE[2]]
    for j in jobs:
      add('role',j['company']+' | '+j['role']);add('meta',j['dates'])
      for text in j[track]:add('bullet',text)
    add('pagebreak','')
    add('title',('Data and BI' if track=='data' else 'Marketing')+' project evidence')
    add('body','Five independent case studies, completed individually from question to recommendation. Fintech expertise is project-based; professional experience is in ecommerce and marketing.')
    for key in t['order']:
      c=CASES[key]
      add('role',c['name'])
      add('meta',(c['domain'] if track=='data' else c['capability'])+' | '+c['duration'])
      for text in CV_PROJECTS[track][key]:add('bullet',text)
      path=('/projects/'+c['data_slug']+'/' if track=='data' else '/work/'+c['marketing_slug']+'/')
      add('case_link',('Read case study',t['site']+path))
    add('heading','Education and development')
    for text in EDUCATION:add('contact',text)
    return out

def footer(canvas,doc):
    canvas.setFont('CVBody',8);canvas.setFillColor(colors.HexColor('#526478'))
    canvas.drawString(40,25,'Phan Hoai Thu')
    canvas.drawRightString(A4[0]-40,25,str(doc.page))

def add_hyperlink(paragraph,label,url):
    rel=paragraph.part.relate_to(url,RT.HYPERLINK,is_external=True)
    link=OxmlElement('w:hyperlink');link.set(qn('r:id'),rel)
    run=OxmlElement('w:r');pr=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),'1254A1');pr.append(color)
    size=OxmlElement('w:sz');size.set(qn('w:val'),'18');pr.append(size);run.append(pr)
    text=OxmlElement('w:t');text.text=label;run.append(text);link.append(run);paragraph._p.append(link)

def make_cv(track,root,track_root=None):
    root=Path(root);base=Path(track_root) if track_root else root/track;out=base/'cv';out.mkdir(parents=True,exist_ok=True)
    name=TRACKS[track]['cv']; story=[];md=[];doc=Document()
    section=doc.sections[0];section.page_width=Inches(8.2677);section.page_height=Inches(11.6929)
    section.top_margin=Inches(.49);section.bottom_margin=Inches(.48);section.left_margin=Inches(.56);section.right_margin=Inches(.56)
    section.header_distance=Inches(.15);section.footer_distance=Inches(.2)
    normal=doc.styles['Normal'];normal.font.name='Arial';normal.font.size=Pt(10.5);normal.font.color.rgb=RGBColor(0,0,0)
    normal.paragraph_format.line_spacing=Pt(14.6);normal.paragraph_format.space_after=Pt(5)
    for st in ['Title','Subtitle','Heading 1','Heading 2']:
      doc.styles[st].font.name='Arial';doc.styles[st].font.color.rgb=RGBColor(0,0,0)
    for border in list(doc.styles.element.iter(qn('w:pBdr'))):
      border.getparent().remove(border)
    for kind,value in blocks(track):
      if kind=='pagebreak':story.append(PageBreak());doc.add_page_break();continue
      if kind in ['links','case_link']:
        links=value if kind=='links' else [value]
        story.append(Paragraph(' | '.join('<link href="'+url+'" color="#1254A1">'+escape(label)+'</link>' for label,url in links),STYLES['contact' if kind=='links' else 'meta']))
        p=doc.add_paragraph();p.paragraph_format.space_after=Pt(3)
        for i,(label,url) in enumerate(links):
          if i:p.add_run(' | ')
          add_hyperlink(p,label,url)
        md.append(' | '.join('['+label+']('+url+')' for label,url in links));continue
      raw='<b>'+escape(value[0])+':</b> '+escape(value[1]) if kind=='skill' else escape(value)
      style='body' if kind=='skill' else kind
      if kind=='bullet':raw='- '+raw
      story.append(Paragraph(raw,STYLES[style]))
      p=doc.add_paragraph(style='Title' if kind=='name' else 'Normal')
      fmt=p.paragraph_format
      if kind in ['name','title','heading','role','meta']:fmt.keep_with_next=True
      fmt.widow_control=True
      if kind=='skill':
        p.add_run(value[0]+': ').bold=True;p.add_run(value[1]);md.append('**'+value[0]+':** '+value[1])
      else:
        run=p.add_run(('- ' if kind=='bullet' else '')+value)
        run.font.size=Pt(STYLES[style].fontSize)
        run.bold=kind in ['name','title','heading','role']
        fmt.line_spacing=Pt(STYLES[style].leading)
        fmt.space_before=Pt(STYLES[style].spaceBefore)
        fmt.space_after=Pt(STYLES[style].spaceAfter)
        if kind=='bullet':fmt.left_indent=Pt(10);fmt.first_line_indent=Pt(-10)
        if kind=='meta':run.font.color.rgb=RGBColor.from_string('415366')
        md.append(('# ' if kind=='name' else '## ' if kind=='heading' else '### ' if kind=='role' else '- ' if kind=='bullet' else '')+value)
    fp=section.footer.paragraphs[0];fp.paragraph_format.space_after=Pt(0)
    r=fp.add_run('Phan Hoai Thu');r.font.size=Pt(8);r.font.color.rgb=RGBColor.from_string('526478')
    doc.core_properties.title=P['name']+' '+TRACKS[track]['role'];doc.core_properties.author=P['name']
    doc.save(out/(name+'.docx'))
    SimpleDocTemplate(str(out/(name+'.pdf')),pagesize=A4,leftMargin=40,rightMargin=40,topMargin=35,bottomMargin=37,title=P['name']+' - '+TRACKS[track]['role'],author=P['name']).build(story,onFirstPage=footer,onLaterPages=footer)
    (out/(name+'.md')).write_text('\n\n'.join(md),encoding='utf-8')
    documents=base/'website'/'documents';documents.mkdir(parents=True,exist_ok=True)
    shutil.copy2(out/(name+'.pdf'),documents/(name+'.pdf'))
    if track=='marketing':
      shutil.copy2(out/(name+'.pdf'),documents/'Phan_Hoai_Thu_Digital_Marketing_CV.pdf')
    print(track,'CV generated')

if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--root',type=Path,required=True);parser.add_argument('--track',choices=TRACKS);args=parser.parse_args()
    for track in ([args.track] if args.track else TRACKS):make_cv(track,args.root)
